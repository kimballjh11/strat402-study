"""STRAT 401 ingest (Phase 0). Run with: npm run ingest:401

Reads inputs/strat401/ (never modifies it) and writes extraction artifacts to work/strat401/:
  export/        the unzipped Canvas export
  text/          every .docx as paragraphs + tables (tables kept as pipe rows)
  images/        images embedded in each .docx
  slides_text/   slide text + speaker notes per deck
  slides_pdf/    decks rendered by LibreOffice
  slides_png/    one PNG per slide (pdftoppm, 60 dpi)
  canvas.json    modules / pages / assignments parsed from viewer/course-data.js
  inventory.md   what exists, what is missing
Renders are cached: a deck is only re-rendered when its PDF is missing.
"""

import html
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
INP = ROOT / "inputs" / "strat401"
WORK = ROOT / "work" / "strat401"
EXPORT = WORK / "export"


IMAGE_NOTES = {
    "image-1.png": "S1 page banner: Porter's generic strategies 2x2 (broad/narrow x lower cost/differentiation)",
    "image.png": "S2 page banner: cover of Minto, The Pyramid Principle",
    "image-3.png": "S3 page banner: generic PEST wheel graphic (Amoeboids), not course-specific content",
    "image-a097aded-2efd-4376-8a6f-18b89f72608d.png": "S4 page: Five Forces diagram (rivalry = 'profit is what remains after the other four take their share')",
    "image-c3bb6d95-4a33-4705-a519-a12bf5010e08.png": "S5 page: HHI C4 Brief Figure 1 - Market A 65/5x7 (C4 80, HHI 4,400, highly concentrated) vs Market B 20x4/5x4 (C4 80, HHI 1,700, moderately concentrated)",
    "image-16cd68f6-5244-455e-99c7-b845c3647bce.png": "S6 page: VRIO gate flowchart (No at V/R/I/O -> disadvantage / parity / temporary / unexploited; all Yes -> sustained)",
    "image-00e1a119-2187-42b0-ba87-53e47e24acb3.png": "VRIO assignment grading rubric (resource identification 20, VRIO application 30, diagnosis 20, recommendations 20, rigor & communication 10)",
    "image-4ffea8e4-023e-4544-ac62-0963815ef8d8.png": "S9 page: value stick (WTP / price / cost / supplier opportunity cost; differentiation raises the top, cost leadership lowers the bottom)",
    "Screenshot 2026-09-23 at 4.09.35\u202fPM.png": "S7 page: Porter generic value chain (Value Chain Brief Exhibit 1)",
}


def log(*a):
    print("[ingest:401]", *a)


def unzip():
    z = INP / "course.zip"
    if not z.exists():
        log("course.zip missing")
        return None
    marker = EXPORT / ".zip-size"
    size = str(z.stat().st_size)
    if not (marker.exists() and marker.read_text() == size):
        if EXPORT.exists():
            shutil.rmtree(EXPORT)
        EXPORT.mkdir(parents=True)
        with zipfile.ZipFile(z) as zf:
            for info in zf.infolist():
                parts = Path(info.filename).parts
                if len(parts) < 2 or info.is_dir():
                    continue
                dest = EXPORT.joinpath(*parts[1:])  # strip the top-level export folder
                if not str(dest.resolve()).startswith(str(EXPORT.resolve())):
                    continue  # zip-slip guard
                dest.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(info) as src, open(dest, "wb") as out:
                    shutil.copyfileobj(src, out)
        marker.write_text(size)
        log("unzipped export")
    return EXPORT / "viewer"


def strip_html(h):
    h = re.sub(r"<(br|/p|/h\d|/li|/tr)[^>]*>", "\n", h or "")
    h = re.sub(r"<li[^>]*>", "  - ", h)
    h = re.sub(r"<img[^>]*src=\"([^\"]+)\"[^>]*>", r"[IMG \1]", h)
    h = re.sub(r"<[^>]+>", "", h)
    return re.sub(r"\n\s*\n+", "\n", html.unescape(h)).strip()


def parse_canvas(viewer):
    s = (viewer / "course-data.js").read_text(encoding="utf-8")
    d = json.loads(s[s.index("{"):].rstrip().rstrip(";"))
    mods = []
    for m in d["modules"]:
        items = []
        for it in m["items"]:
            text = strip_html(it.get("content") or it.get("description") or "")
            kq = re.findall(r"Key Questions?:\s*(.+?)(?:\n(?:Class Materials|\d\.)|$)", text, re.S)
            items.append({
                "title": it.get("title"), "type": it.get("type"), "text": text,
                "dueAt": it.get("dueAt"), "points": it.get("pointsPossible"),
                "keyQuestion": re.sub(r"\s+", " ", kq[0]).strip() if kq else None,
            })
        mods.append({"name": m["name"], "items": items})
    out = {"title": d["title"], "lastDownload": d["lastDownload"], "modules": mods,
           "assignments": [{"title": a["title"], "points": a.get("pointsPossible"), "dueAt": a.get("dueAt")} for a in d["assignments"]]}
    (WORK / "canvas.json").write_text(json.dumps(out, indent=1))
    return out


def extract_docx(src, out_txt, img_dir):
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(src))
    lines = []
    for ch in doc.element.body.iterchildren():
        if ch.tag == qn("w:p"):
            p = Paragraph(ch, doc)
            t = p.text.strip()
            if t:
                st = p.style.name if p.style is not None else ""
                lines.append(("## " if st.startswith("Heading") or st == "Title" else "") + t)
            if p._element.findall(".//" + qn("w:drawing")):
                lines.append("[EMBEDDED IMAGE]")
        elif ch.tag == qn("w:tbl"):
            lines.append("[TABLE]")
            for r in Table(ch, doc).rows:
                cells, prev = [], None
                for c in r.cells:
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(c.text.strip().replace("\n", " / "))
                lines.append(" | ".join(cells))
            lines.append("[/TABLE]")
    out_txt.write_text("\n".join(lines))
    n = 0
    img_dir.mkdir(parents=True, exist_ok=True)
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            n += 1
            (img_dir / f"img{n:02d}{Path(rel.target_ref).suffix}").write_bytes(rel.target_part.blob)
    return len(lines), n


def extract_pptx(src, out_txt):
    from pptx import Presentation

    p = Presentation(str(src))
    rows = []
    for i, s in enumerate(p.slides, 1):
        bits = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = " / ".join(x.strip() for x in sh.text_frame.text.split("\n") if x.strip())
                if t:
                    bits.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    bits.append("TBL: " + " | ".join(c.text.strip() for c in r.cells))
            if sh.shape_type == 13:
                bits.append("[PIC]")
        rows.append(f"--- slide {i}: " + " || ".join(bits))
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            rows.append("    NOTES: " + s.notes_slide.notes_text_frame.text.strip())
    out_txt.write_text("\n".join(rows))
    return len(p.slides)


def render(src):
    pdf_dir, png_root = WORK / "slides_pdf", WORK / "slides_png"
    pdf_dir.mkdir(parents=True, exist_ok=True)
    pdf = pdf_dir / (src.stem + ".pdf")
    if not pdf.exists():
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(pdf_dir), str(src)],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    png = png_root / src.stem
    if pdf.exists() and not (png.exists() and any(png.iterdir())):
        png.mkdir(parents=True, exist_ok=True)
        subprocess.run(["pdftoppm", "-r", "60", "-png", str(pdf), str(png / "s")])


def main():
    WORK.mkdir(parents=True, exist_ok=True)
    viewer = unzip()
    inv = ["# STRAT 401 inventory", "", "Generated by `npm run ingest:401`. Inputs are read-only.", ""]
    if viewer:
        c = parse_canvas(viewer)
        inv += ["## Canvas modules (Key Questions)", ""]
        for m in c["modules"]:
            kq = next((i["keyQuestion"] for i in m["items"] if i["keyQuestion"]), None)
            inv.append(f"- **{m['name'].strip()}**: {kq or '(no key question found)'}")
        inv += ["", "## Assignments", ""]
        inv += [f"- {a['title']} ({a['points']} pts, due {a['dueAt']})" for a in c["assignments"]]

    sources = []
    if viewer:
        sources += sorted((viewer / "files" / "Uploaded Media").glob("*"))
    for sub in ["syllabus", "my-work", "prova", "later"]:
        sources += sorted(p for p in (INP / sub).rglob("*") if p.is_file())

    (WORK / "text").mkdir(exist_ok=True)
    (WORK / "slides_text").mkdir(exist_ok=True)
    inv += ["", "## Files", "", "| File | From | Kind | Extracted |", "|---|---|---|---|"]
    for f in sources:
        frm = "export" if viewer and viewer in f.parents else f.parent.name
        ext = f.suffix.lower()
        tag = f"{frm}__{f.stem}" if frm != "export" else f.stem
        try:
            if ext == ".docx":
                n, imgs = extract_docx(f, WORK / "text" / f"{tag}.txt", WORK / "images" / tag)
                inv.append(f"| {f.name} | {frm} | docx | {n} blocks, {imgs} images |")
            elif ext == ".pptx":
                n = extract_pptx(f, WORK / "slides_text" / f"{tag}.txt")
                render(f)
                inv.append(f"| {f.name} | {frm} | pptx | {n} slides (rendered) |")
            elif ext == ".pdf":
                txt = subprocess.run(["pdftotext", "-layout", str(f), "-"], capture_output=True, text=True).stdout
                (WORK / "text" / f"{tag}.txt").write_text(txt)
                inv.append(f"| {f.name} | {frm} | pdf | {len(txt.split())} words |")
            elif ext in (".png", ".jpg", ".jpeg"):
                inv.append(f"| {f.name} | {frm} | image | view manually |")
            else:
                inv.append(f"| {f.name} | {frm} | {ext or 'other'} | not extracted |")
        except Exception as e:  # keep going; report the failure
            inv.append(f"| {f.name} | {frm} | {ext} | FAILED: {e} |")

    inv += ["", "## Images (identified by viewing)", ""]
    inv += [f"- `{k}`: {v}" for k, v in IMAGE_NOTES.items()]
    inv += ["", "## Missing or empty", ""]
    for sub in ["syllabus", "my-work", "prova", "later"]:
        real = [p for p in (INP / sub).rglob("*") if p.is_file() and p.name != "README.txt"]
        inv.append(f"- `inputs/strat401/{sub}/`: {len(real)} file(s)" + ("" if real else " — EMPTY"))
    inv.append("- Slides for Session 4, 5, 10: none posted (Canvas says so). Session 11 slides: not in export.")
    (WORK / "inventory.md").write_text("\n".join(inv) + "\n")
    log("wrote", WORK / "inventory.md")


if __name__ == "__main__":
    sys.exit(main())
