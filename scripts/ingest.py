"""Phase 0 ingest: extract everything under inputs/ into work/.

Run with: python3 -I scripts/ingest.py  (or `npm run ingest`)
Never modifies inputs/.
"""
import json, os, re, sys, email, quopri, html
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
INP = ROOT / "inputs"
WORK = ROOT / "work"
COURSE = INP / "course" / "viewer"
FILES = COURSE / "files"

from bs4 import BeautifulSoup
import pymupdf
import openpyxl


def slug(s):
    return re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-")[:80]


def html_to_text(h):
    soup = BeautifulSoup(h or "", "html.parser")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for tag in soup.find_all(["p", "li", "h1", "h2", "h3", "h4", "h5", "tr", "div"]):
        tag.insert_after("\n")
    for li in soup.find_all("li"):
        li.insert(0, "- ")
    txt = soup.get_text()
    txt = re.sub(r"[ \t\xa0]+", " ", txt)
    txt = re.sub(r"\n\s*\n+", "\n\n", txt)
    return txt.strip()


inventory = []


def dump_course_data():
    s = (COURSE / "course-data.js").read_text()
    d = json.loads(s[s.index("{"):].rstrip().rstrip(";"))
    out = WORK / "course" / "modules"
    out.mkdir(parents=True, exist_ok=True)
    los = {}
    for m in d["modules"]:
        mdir = out / slug(m["name"])
        mdir.mkdir(exist_ok=True)
        n = 0
        for i, it in enumerate(m["items"]):
            if it.get("content"):
                (mdir / f"{i:02d}-{slug(it['title'])}.txt").write_text(it["title"] + "\n\n" + html_to_text(it["content"]))
                (mdir / f"{i:02d}-{slug(it['title'])}.html").write_text(it["content"])
                n += 1
        inventory.append((f"course-data.js module: {m['name']}", f"{len(m['items'])} items, {n} with content"))
    for key in ["pages", "assignments", "discussion_topics", "quizzes"]:
        kd = WORK / "course" / key
        kd.mkdir(parents=True, exist_ok=True)
        for x in d[key]:
            (kd / f"{slug(x['title'])}.txt").write_text(x["title"] + "\n\n" + html_to_text(x.get("content")) + "\n\nMETA: " + json.dumps({k: v for k, v in x.items() if k != "content"}))
        inventory.append((f"course-data.js {key}", f"{len(d[key])} entries"))
    return d


def dump_pdf(path, rasterize=False):
    doc = pymupdf.open(path)
    name = slug(path.stem)
    tdir = WORK / "text"
    tdir.mkdir(parents=True, exist_ok=True)
    parts = []
    for i, page in enumerate(doc):
        parts.append(f"\n===== PAGE {i+1} =====\n" + page.get_text())
    (tdir / f"{name}.txt").write_text("".join(parts))
    if rasterize:
        rdir = WORK / "slides" / name
        rdir.mkdir(parents=True, exist_ok=True)
        for i, page in enumerate(doc):
            out = rdir / f"p{i+1:02d}.png"
            if not out.exists():
                page.get_pixmap(dpi=110).save(out)
    inventory.append((str(path.relative_to(INP)), f"PDF, {len(doc)} pages" + (", rasterized" if rasterize else "")))


def cell_dump(ws_v, ws_f):
    lines = []
    for row in ws_f.iter_rows():
        for c in row:
            if c.comment is not None:
                lines.append(f"{c.coordinate}\tCOMMENT: {c.comment.text!r}")
            if c.value is None:
                continue
            v = ws_v[c.coordinate].value
            f = c.value
            if isinstance(f, str) and f.startswith("="):
                lines.append(f"{c.coordinate}\t{v!r}\t{f}")
            else:
                lines.append(f"{c.coordinate}\t{v!r}")
    return lines


def dump_xlsx(path, tag):
    try:
        wv = openpyxl.load_workbook(path, data_only=True)
        wf = openpyxl.load_workbook(path, data_only=False)
    except Exception as e:
        inventory.append((str(path.relative_to(INP)), f"XLSX FAILED: {e}"))
        return
    xdir = WORK / "workbooks"
    xdir.mkdir(parents=True, exist_ok=True)
    out = []
    sheets = []
    for name in wf.sheetnames:
        ws_f, ws_v = wf[name], wv[name]
        lines = cell_dump(ws_v, ws_f)
        sheets.append(f"{name} ({ws_f.max_row}x{ws_f.max_column})")
        out.append(f"\n===== SHEET: {name} (dims {ws_f.dimensions}) =====\n" + "\n".join(lines))
    (xdir / f"{tag}--{slug(path.stem)}.txt").write_text("".join(out))
    inventory.append((str(path.relative_to(INP)), "XLSX sheets: " + "; ".join(sheets)))


def dump_pptx(path):
    import zipfile
    z = zipfile.ZipFile(path)
    slides = sorted([n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml", n)], key=lambda n: int(re.findall(r"\d+", n)[0]))
    out = []
    for n in slides:
        x = z.read(n).decode("utf8")
        texts = re.findall(r"<a:t>(.*?)</a:t>", x)
        out.append(f"\n===== SLIDE {re.findall(r'\d+', n)[0]} =====\n" + "\n".join(html.unescape(t) for t in texts))
    notes = sorted([n for n in z.namelist() if re.match(r"ppt/notesSlides/notesSlide\d+\.xml", n)])
    for n in notes:
        x = z.read(n).decode("utf8")
        out.append(f"\n===== NOTES {n} =====\n" + " ".join(html.unescape(t) for t in re.findall(r"<a:t>(.*?)</a:t>", x)))
    mdir = WORK / "slides" / slug(path.stem)
    mdir.mkdir(parents=True, exist_ok=True)
    for n in z.namelist():
        if n.startswith("ppt/media/"):
            (mdir / Path(n).name).write_bytes(z.read(n))
    (WORK / "text" / f"{slug(path.stem)}.txt").write_text("".join(out))
    inventory.append((str(path.relative_to(INP)), f"PPTX, {len(slides)} slides"))


def dump_mhtml(path):
    msg = email.message_from_bytes(path.read_bytes())
    qdir = WORK / "quizzes" / slug(path.stem)
    qdir.mkdir(parents=True, exist_ok=True)
    nimg = 0
    for part in msg.walk():
        ct = part.get_content_type()
        if ct == "text/html" and "results" in part.get("Content-Location", ""):
            (qdir / "page.html").write_bytes(part.get_payload(decode=True))
        elif ct.startswith("image/"):
            loc = part.get("Content-Location", "")
            if "instructure" in loc or "files" in loc or "preview" in loc:
                nimg += 1
                ext = ct.split("/")[1].replace("jpeg", "jpg").split("+")[0]
                (qdir / f"img{nimg:02d}.{ext}").write_bytes(part.get_payload(decode=True))
                (qdir / f"img{nimg:02d}.loc").write_text(loc)
    h = (qdir / "page.html").read_text(errors="ignore")
    (qdir / "page.txt").write_text(html_to_text(h))
    inventory.append((str(path.relative_to(INP)), f"MHTML quiz export, {nimg} images"))


def main():
    dump_course_data()
    for p in sorted(FILES.rglob("*")):
        if p.is_dir() or p.name.startswith("."):
            continue
        rel = p.relative_to(FILES)
        top = rel.parts[0]
        if p.suffix == ".pdf":
            dump_pdf(p, rasterize=(top == "Slides"))
        elif p.suffix == ".xlsx":
            dump_xlsx(p, "course")
        elif p.suffix == ".pptx":
            dump_pptx(p)
        else:
            inventory.append((str(p.relative_to(INP)), f"{p.suffix} file ({p.stat().st_size//1024} KB), not parsed"))
    for sub, tag in [("my-assignments", "mine"), ("later", "later"), ("optional", "optional"), ("textbook", "textbook")]:
        d = INP / sub
        files = [p for p in sorted(d.rglob("*")) if p.is_file() and not p.name.startswith((".", "~$"))]
        if not files:
            inventory.append((f"inputs/{sub}/", "EMPTY"))
        for p in files:
            if p.suffix == ".xlsx":
                dump_xlsx(p, tag)
            elif p.suffix == ".pdf":
                dump_pdf(p, rasterize=(sub == "later"))
            elif p.suffix == ".pptx":
                dump_pptx(p)
            else:
                inventory.append((str(p.relative_to(INP)), "not parsed"))
    qs = sorted((INP / "quizzes").glob("*"))
    if not qs:
        inventory.append(("inputs/quizzes/", "EMPTY"))
    for p in qs:
        if p.suffix == ".mhtml":
            dump_mhtml(p)
    with open(WORK / "inventory.md", "w") as f:
        f.write("# Inventory (generated by scripts/ingest.py)\n\n| File | Contents |\n|---|---|\n")
        for a, b in inventory:
            f.write(f"| {a} | {b} |\n")
    print(f"{len(inventory)} inventory rows")


if __name__ == "__main__":
    main()
